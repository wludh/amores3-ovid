#!/usr/bin/env python3
"""Export reviewed editorial notes to Word without modifying the TEI archive."""
import argparse
from collections import Counter
from datetime import date
import json
import re
from pathlib import Path
import subprocess
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NS = {'tei': 'http://www.tei-c.org/ns/1.0'}
XML_ID = '{http://www.w3.org/XML/1998/namespace}id'
TYPES = {'observation': 'Observation', 'uncertainty': 'Documented uncertainty', 'review': 'Editorial question'}
WITNESSES = ('P', 'Y', 'S', 'O', 'LL')


def collect_notes():
    records = json.loads((ROOT / 'docs/data/tei-corpus.json').read_text())['records']
    result = []
    for record in sorted(records, key=lambda r: (int(r['poem'].split('.')[1]), WITNESSES.index(r['witness']))):
        assert record['status'] == 'reviewed', 'Only reviewed corpus notes may be exported'
        root = ET.parse(ROOT / 'docs' / record['path']).getroot()
        by_id = {el.get(XML_ID): el for el in root.iter() if el.get(XML_ID)}
        for note in root.findall('.//tei:note', NS):
            if note.get('type') not in TYPES:
                continue  # Manuscript inscriptions are source text, not editorial prose.
            target = note.get('target', '').split()[0].removeprefix('#') if note.get('target') else ''
            source = by_id.get(target)
            if source is not None and source.tag == '{%s}gap' % NS['tei'] and source.get('reason') == 'not-transmitted':
                label = 'Source coverage'
            elif 'lower-margin' in target:
                label = 'Lower margin'
            else:
                label = 'Line ' + note.get('n', 'unspecified')
            result.append({
                'poem': record['poem'], 'witness': record['witness'],
                'id': note.get(XML_ID), 'label': label, 'type': note.get('type'),
                'line': note.get('n', ''),
                'target': note.get('target'), 'text': ' '.join(''.join(note.itertext()).split()),
            })
    assert len({n['id'] for n in result}) == len(result), 'Duplicate note IDs'
    return sorted(result, key=lambda n: (
        int(n['poem'].split('.')[1]), WITNESSES.index(n['witness']),
        int(re.search(r'\d+', n['line']).group()) if re.search(r'\d+', n['line']) else 10000,
    ))


def bookmark(paragraph, name, number):
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(number)); start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(number))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def field(paragraph, instruction):
    node = OxmlElement('w:fldSimple'); node.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'; run.append(text); node.append(run)
    paragraph._p.append(node)


def build(output):
    notes = collect_notes()
    doc = Document()
    # Remove Word's inherited decorative title rule from the default template.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    props = doc.core_properties
    props.title = 'Amores III Transcription Notes'
    props.subject = 'Machine generated transcription notes exported from the reviewed TEI corpus'
    props.author = 'OpenAI Codex'
    props.last_modified_by = 'OpenAI Codex'
    props.comments = 'Machine-generated reference material. Not authored or approved by Professor Dance.'
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    section.header_distance = section.footer_distance = Inches(.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(25)
    doc.styles['Title'].paragraph_format.space_after = Pt(10)
    doc.styles['Subtitle'].font.size = Pt(12)
    for name, size in [('Heading 1', 18), ('Heading 2', 13)]:
        style = doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
    doc.styles['Heading 1'].paragraph_format.page_break_before = True
    header = section.header.paragraphs[0]
    header.text = 'AMORES III   |   MACHINE GENERATED TRANSCRIPTION NOTES'
    header.style = doc.styles['Header']; header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run('OpenAI Codex   ·   Page ').font.size = Pt(8)
    field(footer, 'PAGE')
    doc.add_paragraph('Amores III Transcription Notes', 'Title')
    doc.add_paragraph('Machine generated reference material', 'Subtitle')
    doc.add_paragraph(
        'Prepared by OpenAI Codex for James Kull. These notes were generated and reviewed by AI agents during transcription of the Amores III witnesses. They were not written by Professor Dance and do not represent his commentary or scholarly approval.'
    )
    doc.add_paragraph(
        f'This document collects all {len(notes):,} editorial transcription notes from poems 3.1–3.15, organized by poem, witness and line. The note wording and uncertainty are preserved from the reviewed TEI, with XML formatting whitespace normalized. Manuscript inscriptions and alteration tags remain part of the source transcription on the site.'
    )
    counts = Counter(n['witness'] for n in notes)
    doc.add_paragraph('Coverage: ' + '; '.join(f'{w} — {counts[w]:,} notes' for w in WITNESSES) + '. LL has no editorial transcription notes, so it has no repeated empty sections.')
    doc.add_paragraph('Observation records a supported feature or reading. Documented uncertainty retains an unresolved image reading or interpretation without assigning a user task. The exported corpus contains no open editorial questions.')
    commit = subprocess.check_output(['git', 'rev-parse', '--short', 'HEAD'], cwd=ROOT, text=True).strip()
    doc.add_paragraph(f'Exported {date.today():%d %B %Y} from the reviewed corpus at repository commit {commit}. Source archive: docs/data/tei. Review evidence: scripts/output.').runs[0].font.size = Pt(9)
    doc.add_paragraph('Finding a note', 'Heading 2')
    doc.add_paragraph('Use Word’s Navigation Pane to select a poem or witness. Within each witness, the bold label identifies the verse or source-coverage note. The following links jump to each poem.')
    for n in range(1, 16):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), f'poem_{n}')
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr'); color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
        run.append(props); text = OxmlElement('w:t'); text.text = f'Amores 3.{n}'; run.append(text); link.append(run); p._p.append(link)
    for n in range(1, 16):
        poem = f'3.{n}'
        heading = doc.add_paragraph(f'Amores III Poem {n}', 'Heading 1'); bookmark(heading, f'poem_{n}', n)
        for witness in WITNESSES:
            selected = [note for note in notes if note['poem'] == poem and note['witness'] == witness]
            if not selected:
                continue
            doc.add_paragraph(f'Witness {witness}', 'Heading 2')
            for note in selected:
                p = doc.add_paragraph()
                p.paragraph_format.keep_together = True
                p.add_run(f'{note["label"]} · {TYPES[note["type"]]}. ').bold = True
                p.add_run(note['text'])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    # Every source note must be present exactly once, in the exported order.
    exported = Document(output)
    actual = [p.text for p in exported.paragraphs if p.runs and p.runs[0].bold]
    expected = [f'{n["label"]} · {TYPES[n["type"]]}. {n["text"]}' for n in notes]
    assert actual == expected, 'Export changed, omitted or reordered a note'
    print(json.dumps({'output': str(output), 'notes': len(notes), 'by_witness': dict(counts), 'verified': True}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'exports/Amores III Transcription Notes.docx')
    build(parser.parse_args().output.resolve())
